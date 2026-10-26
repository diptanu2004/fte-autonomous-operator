"""Shared-workspace file access (sandboxed to one root) and the email outbox."""
from __future__ import annotations

import csv
import io
import re
import time
from pathlib import Path


class FileError(RuntimeError):
    pass


class Workspace:
    def __init__(self, root: Path, max_chars: int = 8000):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.max_chars = max_chars

    def resolve(self, rel: str) -> Path:
        p = (self.root / str(rel).replace("\\", "/").lstrip("/")).resolve()
        if p != self.root and self.root not in p.parents:
            raise FileError(f"Path '{rel}' is outside the workspace.")
        return p

    def rel(self, p: Path) -> str:
        return p.resolve().relative_to(self.root).as_posix()

    def list(self, rel: str = "") -> str:
        base = self.resolve(rel)
        if not base.exists():
            raise FileError(f"'{rel}' does not exist.")
        items = [p for p in sorted(base.rglob("*")) if p.is_file()]
        if not items:
            return f"(no files under '{rel or '.'}')"
        return "\n".join(f"{self.rel(p)}  ({p.stat().st_size} bytes)" for p in items[:200])

    def read(self, rel: str) -> str:
        p = self.resolve(rel)
        if not p.is_file():
            raise FileError(f"File '{rel}' not found. Use list_files to see what exists.")
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            text = self._xlsx(p)
        elif suffix == ".pdf":
            text = self._pdf(p)
        elif suffix == ".csv":
            rows = list(csv.reader(io.StringIO(p.read_text(encoding="utf-8-sig"))))
            text = "\n".join(" | ".join(r) for r in rows) + self._column_totals(rows)
        else:
            text = p.read_text(encoding="utf-8", errors="replace")
        if len(text) > self.max_chars:
            text = text[: self.max_chars] + f"\n...[truncated {len(text) - self.max_chars} chars]"
        return text

    @staticmethod
    def _column_totals(rows: list[list[str]]) -> str:
        """Sum amount-like columns of line-item tables so totals are never added up "by eye"."""
        if len(rows) < 3:  # header + at least two data rows
            return ""
        header, data = rows[0], rows[1:]
        sums = []
        for i, name in enumerate(header):
            if not re.search(r"(?i)total|amount|sum|value", name):
                continue
            vals = [r[i].replace(",", "").replace("$", "").strip() for r in data if i < len(r) and r[i].strip()]
            try:
                total = sum(float(v) for v in vals)
            except ValueError:
                continue
            sums.append(f"{name} = {total:.2f} (sum of {len(vals)} rows)")
        return ("\nCOMPUTED COLUMN TOTALS: " + "; ".join(sums)) if sums else ""

    @staticmethod
    def _xlsx(p: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    out.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out)

    @staticmethod
    def _pdf(p: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise FileError("PDF support needs `pip install pypdf`.") from e
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)

    def write(self, rel: str, content: str) -> Path:
        p = self.resolve(rel)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
        return p

    def send_email(self, to: str, subject: str, body: str, sender: str) -> Path:
        """Delivery is simulated: the message lands in outbox/ as the system of record."""
        if not to or not subject or not body:
            raise FileError("send_email needs to, subject and body.")
        for addr in re.split(r"[,;]", to):
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", addr.strip()):
                raise FileError(f"'{addr.strip()}' is not a valid email address.")
        slug = re.sub(r"[^a-z0-9]+", "-", subject.lower()).strip("-")[:50]
        stamp = time.strftime("%Y%m%d-%H%M%S")
        msg = f"From: {sender}\nTo: {to}\nDate: {time.strftime('%a, %d %b %Y %H:%M:%S')}\nSubject: {subject}\n\n{body}\n"
        return self.write(f"outbox/{stamp}-{slug}.eml", msg)
